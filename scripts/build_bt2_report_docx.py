"""Generate BT2 Word report from project docs."""
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

OUT_PATHS = [
    Path(r"d:\cdtn1-nguyenphamquangthien-l2\docs\BT2_2374802011250_NguyenPhamQuangThien.docx"),
    Path(r"C:\Users\Asus\Downloads\BT2_2374802011250_NguyenPhamQuangThien.docx"),
]


def set_run_font(run, size=13, bold=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold


def add_center(doc, text, size=13, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size, bold)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, 14 if level == 1 else 13, True)


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)


def add_screenshot_placeholder(doc, code, title, caption):
    add_para(doc, f"[{code} — {title}]")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("(Chèn ảnh tại đây — xem docs/screenshot-checklist.md)")
    set_run_font(run, 11)
    run.italic = True
    add_para(doc, f"Hình {code.split()[1]}: {caption}")
    doc.add_paragraph()


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Trang bìa
    add_center(doc, "TRƯỜNG ĐẠI HỌC VĂN LANG", 13, True)
    add_center(doc, "KHOA CÔNG NGHỆ THÔNG TIN", 13, True)
    doc.add_paragraph()
    add_center(doc, "Môn: Chuyên đề Tốt nghiệp 1 (71ITGR40203)", 12)
    add_center(doc, "BÀI TẬP 2: HIỆN THỰC HÓA MODULE CỐT LÕI", 14, True)
    add_center(doc, "Đề tài: Mekong Mobile Smart CRM", 12)
    add_center(doc, "Luồng nghiệp vụ: L2 – Tiếp nhận và phân loại yêu cầu bảo hành", 12)
    add_center(doc, "Track: Software Engineering (SE)", 12)
    doc.add_paragraph()
    add_center(doc, "Sinh viên: Nguyễn Phạm Quang Thiện", 12)
    add_center(doc, "MSSV: 2374802011250", 12)
    add_center(doc, "GitHub: https://github.com/Thien1707/cdtn1-nguyenphamquangthien-l2", 11)
    add_center(doc, "Tag nộp: v0.2", 11)
    doc.add_paragraph()
    add_center(doc, "TP. Hồ Chí Minh – Năm 2026", 12)
    doc.add_page_break()

    add_heading(doc, "MỤC LỤC (gợi ý)", 1)
    add_bullets(
        doc,
        [
            "1. Giới thiệu và mục tiêu BT2",
            "2. Kiến trúc và công nghệ",
            "3. Module cốt lõi và quy tắc nghiệp vụ",
            "4. Module mở rộng và giao diện",
            "5. Kiểm thử",
            "6. Hướng dẫn cài đặt và tái lập (README)",
            "7. Thay đổi so với BT1",
            "8. Minh chứng hình ảnh",
            "9. Rubric tự đánh giá",
            "10. Khai báo AI",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "1. Giới thiệu và mục tiêu BT2", 1)
    add_para(
        doc,
        "Bài tập 2 hiện thực hóa thiết kế Bài tập 1 thành hệ thống chạy được, có tách lớp, "
        "có kiểm thử và có thể tái lập theo README. Dự án Mekong Mobile Smart CRM (luồng L2) "
        "hỗ trợ tiếp nhận yêu cầu bảo hành, tra cứu khách hàng và thiết bị, tính SLA theo giờ làm việc, "
        "phân công kỹ thuật viên và báo cáo phiếu quá hạn cam kết."
    )
    add_para(
        doc,
        "Sản phẩm nộp gồm: mã nguồn trên GitHub (tag v0.2), tài liệu trong thư mục docs/ "
        "và báo cáo PDF này."
    )

    add_heading(doc, "2. Kiến trúc và công nghệ", 1)
    add_para(
        doc,
        "Kiến trúc 3 lớp: Routes → Controllers → Services → Repositories → PostgreSQL "
        "(hoặc bộ nhớ in-memory khi STORAGE_DRIVER=memory). Frontend tĩnh phục vụ tại src/public/, "
        "gọi REST API qua lớp client src/public/api/client.js."
    )
    add_screenshot_placeholder(
        doc,
        "SCREENSHOT S11",
        "Sơ đồ kiến trúc thực tế",
        "Kiến trúc phân lớp đã hiện thực (Routes, Controllers, Services, Repositories, DB, API client). "
        "Nguồn: docs/architecture-thuc-te.md.",
    )
    add_screenshot_placeholder(
        doc,
        "SCREENSHOT S10",
        "Cấu trúc ba lớp backend",
        "Minh họa thư mục src/ và file service không phụ thuộc HTTP.",
    )

    add_heading(doc, "3. Module cốt lõi và quy tắc nghiệp vụ", 1)
    add_heading(doc, "3.1. Hai thực thể CRUD", 2)
    add_para(
        doc,
        "Đáp ứng rubric S1: thực thể customer và device có đủ Create, Read, Update, Delete. "
        "Phiếu bảo hành (ticket) không có thao tác xóa tùy tiện theo nghiệp vụ BT1."
    )
    add_bullets(
        doc,
        [
            "Customer: GET/POST /api/v1/customers, GET/PATCH/DELETE /api/v1/customers/:id, GET search theo SĐT.",
            "Device: POST /api/v1/devices, GET/PATCH/DELETE /api/v1/devices/:id.",
            "Kiểm thử: tests/crud.api.test.js.",
        ],
    )
    add_screenshot_placeholder(doc, "SCREENSHOT S02", "CRUD khách hàng", "Tạo, đọc, sửa, xóa khách hàng trên giao diện hoặc API.")
    add_screenshot_placeholder(doc, "SCREENSHOT S03", "CRUD thiết bị", "Tạo, đọc, sửa, xóa thiết bị.")

    add_heading(doc, "3.2. Quy tắc nghiệp vụ", 2)
    add_bullets(
        doc,
        [
            "QT-01: Số điện thoại đúng 10 chữ số, không trùng khi tạo khách.",
            "QT-03: Mô tả lỗi bắt buộc, tối thiểu 10 ký tự sau trim.",
            "QT-04: SLA theo ưu tiên CAO 24h, TRUNG_BINH 72h, THAP 120h giờ làm việc; T2–T7 08:00–17:00 (UTC+7), không tính CN.",
            "Phân công kỹ thuật viên: level ≥ 3, cùng trung tâm phiếu.",
            "Kỹ thuật viên: số điện thoại khách che dạng 090****567 trong API queue.",
            "Tạo phiếu: thiết bị phải thuộc khách hàng đã chọn; ghi log trạng thái MOI khi tiếp nhận.",
        ],
    )
    add_screenshot_placeholder(doc, "SCREENSHOT S04", "Validation dữ liệu sai", "Thông báo lỗi SĐT hoặc mô tả lỗi dưới 10 ký tự.")
    add_screenshot_placeholder(doc, "SCREENSHOT S05", "Tính SLA tự động", "Cột hạn SLA sau khi tạo phiếu.")

    add_heading(doc, "4. Module mở rộng và giao diện", 1)
    add_para(
        doc,
        "Ít nhất ba màn hình gọi API thật: Tiếp nhận phiếu, Danh sách phiếu bảo hành, "
        "Phân công / Phiếu được giao / Báo cáo quá hạn (theo vai trò). "
        "Màn hình xử lý trạng thái đang tải, rỗng (kể cả bộ lọc không ra kết quả) và lỗi có nút Thử lại."
    )
    add_screenshot_placeholder(doc, "SCREENSHOT S01", "Tiếp nhận phiếu", "Luồng tra cứu khách, chọn thiết bị, tạo phiếu và hiển thị mã BH.")
    add_screenshot_placeholder(doc, "SCREENSHOT S06", "Phân công kỹ thuật viên", "Gán phiếu cho KTV đủ level.")
    add_screenshot_placeholder(doc, "SCREENSHOT S07", "Ba trạng thái UI", "Loading, empty, error và retry khi tắt backend.")

    add_heading(doc, "5. Kiểm thử", 1)
    add_para(doc, "Lệnh chạy test: npm test (một lệnh duy nhất, ghi trong README).")
    add_para(doc, "Kết quả thực thi (cập nhật khi nộp): 17 test PASS, 0 FAIL, 1 skipped (PostgreSQL integration nếu thiếu TEST_DATABASE_URL).")
    add_bullets(
        doc,
        [
            "Loại 1 — Luồng đúng: CRUD, tạo phiếu, tra cứu, phân công, báo cáo.",
            "Loại 2 — Biên: SLA qua cuối tuần, Chủ nhật; mô tả 9 vs 10 ký tự.",
            "Loại 3 — Dữ liệu sai: SĐT sai, priority invalid, level KTV thấp, thiếu lý do CHO_LINH_KIEN.",
        ],
    )
    add_screenshot_placeholder(doc, "SCREENSHOT S08", "Kết quả npm test", "Terminal hiển thị số test pass.")

    add_heading(doc, "6. Hướng dẫn cài đặt và tái lập", 1)
    add_heading(doc, "6.1. Yêu cầu môi trường", 2)
    add_bullets(doc, ["Node.js 18+ hoặc 20+", "npm 9+", "PostgreSQL 14+ (đường chạy chính thức)", "PowerShell hoặc bash"])

    add_heading(doc, "6.2. Các bước chạy (PostgreSQL)", 2)
    steps = [
        "npm install",
        "copy src\\.env.example src\\.env — sửa DATABASE_URL và JWT_SECRET",
        "Tạo database rỗng mekong_crm",
        "npm run db:init",
        "Tạo tài khoản manager: npm run user:create (xem README)",
        "npm run dev",
        "Mở http://localhost:3000 và đăng nhập",
    ]
    for i, step in enumerate(steps, 1):
        add_para(doc, f"Bước {i}. {step}")

    add_para(doc, "Chạy lại seed (C5): npm run db:init lần 2 — không nhân đôi dữ liệu (ON CONFLICT DO NOTHING).")
    add_screenshot_placeholder(
        doc,
        "SCREENSHOT S09",
        "Migration và seed",
        "Kết quả db:init lần hai, số dòng customer không tăng.",
    )
    add_para(
        doc,
        "Demo không PostgreSQL: đặt STORAGE_DRIVER=memory; tài khoản frontdesk/manager/technician, mật khẩu Welcome123!.",
    )

    add_heading(doc, "6.3. Xác minh hệ thống chạy đúng", 2)
    add_bullets(
        doc,
        [
            "Tra cứu 0901234567, tạo phiếu mô tả ≥10 ký tự.",
            "Manager phân công KTV level ≥3.",
            "Technician xem queue, SĐT đã che.",
            "npm test pass hết.",
        ],
    )

    add_heading(doc, "7. Thay đổi so với BT1", 1)
    add_para(doc, "Chi tiết đầy đủ: docs/thay-doi-thiet-ke.md (TD-01 … TD-12). Tóm tắt BT2:")
    add_bullets(
        doc,
        [
            "TD-07: CRUD customer và device.",
            "TD-08: Tách API client frontend.",
            "TD-09: Mô tả lỗi tối thiểu 10 ký tự.",
            "TD-10: Giờ làm việc SLA T2–T7 08:00–17:00.",
            "TD-11: Log MOI khi tạo phiếu.",
            "TD-12: Tài liệu kiến trúc thực tế.",
        ],
    )

    add_heading(doc, "8. Bảng checklist ảnh minh chứng", 1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Mã"
    hdr[1].text = "Tên"
    hdr[2].text = "File đề xuất"
    hdr[3].text = "Trạng thái"
    rows = [
        ("S01", "Tiếp nhận phiếu", "docs/screenshots/S01-tiep-nhan.png", "Chưa chụp"),
        ("S02", "CRUD customer", "S02-crud-customer.png", "Chưa chụp"),
        ("S03", "CRUD device", "S03-crud-device.png", "Chưa chụp"),
        ("S04", "Validation", "S04-validation.png", "Chưa chụp"),
        ("S05", "SLA", "S05-sla.png", "Chưa chụp"),
        ("S06", "Phân công", "S06-assign.png", "Chưa chụp"),
        ("S07", "UI states", "S07-ui-states.png", "Chưa chụp"),
        ("S08", "npm test", "S08-npm-test.png", "Chưa chụp"),
        ("S09", "Seed idempotent", "S09-seed-idempotent.png", "Chưa chụp"),
        ("S10", "Ba lớp", "S10-layering.png", "Chưa chụp"),
        ("S11", "Kiến trúc", "architecture-thuc-te.png", "Chưa chụp"),
        ("S12", "README (tuỳ chọn)", "S12-readme.png", "Chưa chụp"),
    ]
    for code, name, path, status in rows:
        row = table.add_row().cells
        row[0].text = code
        row[1].text = name
        row[2].text = path
        row[3].text = status

    doc.add_paragraph()
    add_heading(doc, "9. Rubric tự đánh giá BT2 (SE)", 1)
    rubric = doc.add_table(rows=1, cols=3)
    rubric.rows[0].cells[0].text = "Mục"
    rubric.rows[0].cells[1].text = "Trạng thái"
    rubric.rows[0].cells[2].text = "Bằng chứng"
    rubric_data = [
        ("C1 Chạy theo README", "Chưa xác minh clone sạch", "Sinh viên chạy phép thử README"),
        ("C2 Tách lớp", "Đạt", "routes → controllers → services → repositories"),
        ("C3 ≥2 quy tắc NV", "Đạt", "QT-01, QT-03, QT-04, level KTV, che SĐT"),
        ("C4 ≥3 test 3 loại", "Đạt", "npm test 17 pass"),
        ("C5 Chạy lại init", "Chưa xác minh", "S09, ON CONFLICT seed"),
        ("C6 .env.example", "Đạt", "src/.env.example"),
        ("C7 Input sai", "Đạt", "HTTP 400, message rõ"),
        ("S1 CRUD 2 entity", "Đạt", "tests/crud.api.test.js"),
        ("S2 Service sạch HTTP", "Đạt", "ticket.service.js"),
        ("S3 Migration + seed", "Đạt", "schema.sql, seed.sql, db:init"),
        ("S4 UI + client + 3 state", "Đạt một phần", "client.js, ui-state — cần S07"),
        ("Tag v0.2", "Chưa", "Push tag sau khi hoàn tất"),
        ("Ảnh minh chứng", "Chưa", "Mục 8"),
    ]
    for a, b, c in rubric_data:
        r = rubric.add_row().cells
        r[0].text = a
        r[1].text = b
        r[2].text = c

    doc.add_paragraph()
    add_heading(doc, "10. Khai báo sử dụng công cụ AI", 1)
    add_para(
        doc,
        "Nội dung đầy đủ và cột “Em kiểm chứng mã AI thế nào” tại docs/ai-disclosure.md trong repository. "
        "Sinh viên chịu trách nhiệm rà soát và ký cam kết trước khi nộp."
    )

    add_heading(doc, "Phụ lục — API tiêu biểu", 1)
    api = doc.add_table(rows=1, cols=3)
    api.rows[0].cells[0].text = "Method"
    api.rows[0].cells[1].text = "URL"
    api.rows[0].cells[2].text = "Mô tả"
    apis = [
        ("POST", "/api/v1/auth/login", "JWT"),
        ("GET", "/api/v1/customers/search?phone=", "Tra cứu QT-01"),
        ("CRUD", "/api/v1/customers, /customers/:id", "Khách hàng"),
        ("CRUD", "/api/v1/devices/:id", "Thiết bị"),
        ("POST", "/api/v1/tickets", "Tạo phiếu + SLA"),
        ("POST", "/api/v1/tickets/:id/assign", "Phân công level ≥3"),
        ("GET", "/api/v1/reports/overdue-tickets", "Báo cáo quá hạn"),
    ]
    for m, u, d in apis:
        row = api.add_row().cells
        row[0].text = m
        row[1].text = u
        row[2].text = d

    add_para(doc, "")
    add_center(doc, "— Hết báo cáo —", 12)

    for path in OUT_PATHS:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
        print("Saved:", path)


if __name__ == "__main__":
    build()
